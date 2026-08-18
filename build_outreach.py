"""Add outreach tables (appointments, contacts, diagnoses) to the existing coverage.db,
without touching the billing tables build_db.py owns (items/covered_daily/summary).

Run whenever the retention workbook or the diagnosis CSV refreshes:
    python3 build_outreach.py /path/to/Retention_Master.xlsx /path/to/NewPatientLook.csv
Defaults to the Downloads copies if no paths are given.
"""
import csv
import sqlite3
import sys
from pathlib import Path

import openpyxl

DEFAULT_RETENTION = Path.home() / "Downloads" / "Retention_Master (2).xlsx"
DEFAULT_DIAGNOSIS = Path.home() / "Downloads" / "Metabolic.Health Diagnosis - NewPatientLook (1).csv"
DB_PATH = Path(__file__).parent / "data" / "coverage.db"


def normalize_mrn(raw):
    digits = "".join(ch for ch in str(raw) if ch.isdigit())
    return digits.lstrip("0") or "0"


def build(retention_path, diagnosis_path):
    if not DB_PATH.exists():
        raise SystemExit(f"{DB_PATH} not found — run build_db.py first.")

    con = sqlite3.connect(DB_PATH)
    cur = con.cursor()

    cur.executescript("""
        DROP TABLE IF EXISTS appointments;
        DROP TABLE IF EXISTS contacts;
        DROP TABLE IF EXISTS diagnoses;
        CREATE TABLE appointments(
            mrn TEXT, appointment_date TEXT, attended INTEGER,
            disease_category TEXT, care_package TEXT
        );
        CREATE TABLE contacts(
            mrn TEXT PRIMARY KEY, first_name TEXT, surname TEXT,
            mobile_phone TEXT, email TEXT, last_visit_date TEXT,
            churn_score REAL, churn_risk TEXT
        );
        CREATE TABLE diagnoses(
            mrn TEXT PRIMARY KEY, diagnosis_groups TEXT
        );
    """)

    wb = openpyxl.load_workbook(retention_path, data_only=True)

    n_appt = 0
    ws = wb["appointment_booked+attended"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        mrn_raw = row[0]
        if not mrn_raw:
            continue
        appt_date = row[2]
        if appt_date is None:
            continue
        date_str = appt_date.date().isoformat() if hasattr(appt_date, "date") else str(appt_date)[:10]
        attended = 1 if row[8] else 0
        cur.execute(
            "INSERT INTO appointments VALUES (?,?,?,?,?)",
            (normalize_mrn(mrn_raw), date_str, attended, row[10], row[9]),
        )
        n_appt += 1

    n_contacts = 0
    ws = wb["Retention_List"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        mrn_raw = row[0]
        if not mrn_raw:
            continue
        last_visit = row[12]
        last_visit_str = last_visit.date().isoformat() if hasattr(last_visit, "date") else (str(last_visit)[:10] if last_visit else None)
        cur.execute(
            "INSERT OR REPLACE INTO contacts VALUES (?,?,?,?,?,?,?,?)",
            (normalize_mrn(mrn_raw), row[1], row[2], row[7], row[6],
             last_visit_str, row[17], row[18]),
        )
        n_contacts += 1

    n_diag = 0
    with open(diagnosis_path, newline="", encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            mrn_raw = row["HospitalNumber"]
            if not mrn_raw:
                continue
            groups = row["Diagnosis Groups"].strip()
            cur.execute(
                "INSERT OR REPLACE INTO diagnoses VALUES (?,?)",
                (normalize_mrn(mrn_raw), groups),
            )
            n_diag += 1

    cur.execute("CREATE INDEX idx_appt_mrn ON appointments(mrn, appointment_date)")
    cur.execute("CREATE INDEX idx_appt_date ON appointments(appointment_date, attended)")

    con.commit()
    con.close()
    print(f"Loaded into {DB_PATH}")
    print(f"  appointments: {n_appt}")
    print(f"  contacts: {n_contacts}")
    print(f"  diagnoses: {n_diag}")


if __name__ == "__main__":
    args = sys.argv[1:]
    retention = Path(args[0]) if len(args) > 0 else DEFAULT_RETENTION
    diagnosis = Path(args[1]) if len(args) > 1 else DEFAULT_DIAGNOSIS
    build(retention, diagnosis)
